"""
NoteWise Dashboard
Main Streamlit application with enhanced export capabilities
"""

import streamlit as st
import os
import sys
import tempfile

# Add project root to path
parent_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if parent_dir not in sys.path:
    sys.path.insert(0, parent_dir)

# Import Gemini extractor
try:
    from core.ocr.gemini_extractor import GeminiExtractor
    GEMINI_EXTRACTION_AVAILABLE = True
except ImportError as e:
    print(f"Gemini extraction not available: {e}")
    GeminiExtractor = None
    GEMINI_EXTRACTION_AVAILABLE = False

# Import Gemini modules
try:
    from core.summarizer.gemini_summarizer import GeminiTextSummarizer
    from core.question_gen.gemini_generator import GeminiQuestionGenerator
    GEMINI_AVAILABLE = True
except ImportError as e:
    print(f"Gemini modules not available: {e}")
    GeminiTextSummarizer = None
    GeminiQuestionGenerator = None
    GEMINI_AVAILABLE = False

# Load environment variables
from dotenv import load_dotenv
load_dotenv()

# Storage system
try:
    from core.storage import SimpleStorage
    STORAGE_AVAILABLE = True
except ImportError as e:
    print(f"Storage not available: {e}")
    SimpleStorage = None
    STORAGE_AVAILABLE = False

# Enhanced export system
try:
    from core.exports import EnhancedExporter
    EXPORT_AVAILABLE = True
except ImportError as e:
    print(f"Enhanced export not available: {e}")
    EnhancedExporter = None
    EXPORT_AVAILABLE = False

def main():
    """Main dashboard function"""
    
    # Initialize storage system
    if STORAGE_AVAILABLE:
        if 'storage' not in st.session_state:
            st.session_state.storage = SimpleStorage()
        if 'session_id' not in st.session_state:
            st.session_state.session_id = st.session_state.storage.create_session()
    
    # Page header
    st.title("🧠 NoteWise - Smart Study Assistant")
    st.markdown("Transform your notes into interactive study material with AI")
    
    # Sidebar
    with st.sidebar:
        st.header("🎛️ Controls")
        
        # Navigation
        page = st.selectbox(
            "Choose a function:",
            ["📄 Text Extraction", "📝 Smart Summary", "❓ Practice Questions", "📊 History"]
        )
        
        # Gemini API Key
        st.markdown("### 🔑 API Configuration")
        gemini_api_key = st.text_input(
            "Gemini API Key:",
            type="password",
            value=os.getenv("GEMINI_API_KEY", ""),
            help="Get your key from https://makersuite.google.com/app/apikey"
        )
        
        # Summary settings
        if page == "📝 Smart Summary":
            st.markdown("### ⚙️ Summary Settings")
            summary_length = st.selectbox("Summary Length:", ["Short", "Standard", "Long"])
        else:
            summary_length = "Standard"
        
        # Question settings
        if page == "❓ Practice Questions":
            st.markdown("### ⚙️ Question Settings")
            question_count = st.slider("Number of Questions:", 3, 15, 5)
        else:
            question_count = 5
        
        # Sidebar branding
        st.markdown("---")
        st.markdown(
            """
            <div style='text-align: center; margin-top: 2rem; padding: 1rem; background: linear-gradient(135deg, #667eea 0%, #764ba2 100%); border-radius: 10px; color: white;'>
                <div style='font-size: 1.1em; font-weight: bold; margin-bottom: 0.5rem;'>
                    🧠 NoteWise
                </div>
                <div style='font-size: 0.8em; opacity: 0.9;'>
                    Smart Study Assistant
                </div>
            </div>
            """, 
            unsafe_allow_html=True
        )
    
    # Initialize session state
    if 'extracted_text' not in st.session_state:
        st.session_state.extracted_text = ""
    if 'summary' not in st.session_state:
        st.session_state.summary = ""
    if 'questions' not in st.session_state:
        st.session_state.questions = []
    
    # Page routing
    if page == "📄 Text Extraction":
        text_extraction_page(gemini_api_key)
    elif page == "📝 Smart Summary":
        summary_page(summary_length, gemini_api_key)
    elif page == "❓ Practice Questions":
        question_generation_page(question_count, gemini_api_key)
    elif page == "📊 History":
        history_page()

def text_extraction_page(gemini_api_key):
    """Page for text extraction from files"""
    st.header("📄 Text Extraction")
    
    # File upload
    uploaded_file = st.file_uploader(
        "Choose a file to extract text from:",
        type=['pdf', 'txt', 'png', 'jpg', 'jpeg', 'pptx', 'ppt'],
        help="Upload PDF, text, image, or PowerPoint files"
    )
    
    if uploaded_file:
        file_type = uploaded_file.type
        
        if file_type == "application/pdf":
            # PDF processing
            with st.spinner("📖 Extracting text from PDF..."):
                try:
                    import pdfplumber
                    
                    # Save uploaded file temporarily
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_file:
                        tmp_file.write(uploaded_file.read())
                        tmp_path = tmp_file.name
                    
                    # Extract text
                    text = ""
                    with pdfplumber.open(tmp_path) as pdf:
                        for page in pdf.pages:
                            page_text = page.extract_text()
                            if page_text:
                                text += page_text + "\n"
                    
                    # Clean up
                    os.unlink(tmp_path)
                    
                    if text.strip():
                        st.session_state.extracted_text = text
                        st.success(f"✅ Extracted {len(text.split())} words from PDF")
                    else:
                        st.error("❌ No text found in PDF")
                        
                except Exception as e:
                    st.error(f"❌ Error processing PDF: {str(e)}")
        
        elif file_type.startswith('image/'):
            # Gemini Vision processing for images
            if GEMINI_EXTRACTION_AVAILABLE and gemini_api_key:
                with st.spinner("🤖 Analyzing image with Gemini Vision..."):
                    try:
                        # Save uploaded image temporarily
                        with tempfile.NamedTemporaryFile(delete=False, suffix=".jpg") as tmp_file:
                            tmp_file.write(uploaded_file.read())
                            tmp_path = tmp_file.name
                        
                        gemini_extractor = GeminiExtractor(api_key=gemini_api_key)
                        result = gemini_extractor.extract_with_analysis(tmp_path)
                        
                        # Clean up
                        os.unlink(tmp_path)
                        
                        if result["extracted_text"].strip():
                            st.session_state.extracted_text = result["extracted_text"]
                            st.success(f"✅ Extracted {result['word_count']} words with Gemini Vision")
                            
                            # Show analysis if available
                            if result.get("analysis"):
                                with st.expander("🔍 Content Analysis"):
                                    st.write(result["analysis"])
                        else:
                            st.error("❌ No text found in image")
                            
                    except Exception as e:
                        st.error(f"❌ Error processing image: {str(e)}")
            else:
                st.error("❌ Gemini Vision not available. Please check your API key.")
        
        elif file_type in ["application/vnd.openxmlformats-officedocument.presentationml.presentation", 
                          "application/vnd.ms-powerpoint"]:
            # PowerPoint processing
            with st.spinner("📊 Extracting text from PowerPoint..."):
                try:
                    from pptx import Presentation
                    
                    # Save uploaded file temporarily
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx") as tmp_file:
                        tmp_file.write(uploaded_file.read())
                        tmp_path = tmp_file.name
                    
                    # Extract text from all slides
                    text = ""
                    presentation = Presentation(tmp_path)
                    
                    for slide_num, slide in enumerate(presentation.slides, 1):
                        text += f"\n--- Slide {slide_num} ---\n"
                        
                        # Extract text from all shapes
                        for shape in slide.shapes:
                            if hasattr(shape, "text") and shape.text.strip():
                                text += shape.text + "\n"
                        
                        # Extract text from notes (if any)
                        if slide.notes_slide and slide.notes_slide.notes_text_frame:
                            notes_text = slide.notes_slide.notes_text_frame.text.strip()
                            if notes_text:
                                text += f"Notes: {notes_text}\n"
                    
                    # Clean up
                    os.unlink(tmp_path)
                    
                    if text.strip():
                        st.session_state.extracted_text = text
                        slide_count = len(presentation.slides)
                        word_count = len(text.split())
                        st.success(f"✅ Extracted {word_count} words from {slide_count} slides")
                        
                        # Show slide breakdown
                        with st.expander(f"📊 Slide Breakdown ({slide_count} slides)"):
                            st.write(f"**Total words:** {word_count}")
                            st.write(f"**Average words per slide:** {word_count // slide_count if slide_count > 0 else 0}")
                    else:
                        st.error("❌ No text found in PowerPoint presentation")
                        
                except ImportError:
                    st.error("❌ PowerPoint support not installed. Please install python-pptx package.")
                except Exception as e:
                    st.error(f"❌ Error processing PowerPoint: {str(e)}")
        
        elif file_type == "text/plain":
            # Text file processing
            try:
                text = uploaded_file.read().decode('utf-8')
                st.session_state.extracted_text = text
                st.success(f"✅ Loaded {len(text.split())} words from text file")
            except Exception as e:
                st.error(f"❌ Error reading text file: {str(e)}")
    
    # Manual text input
    st.markdown("### ✍️ Or enter text manually:")
    manual_text = st.text_area(
        "Enter your text here:",
        height=200,
        placeholder="Paste or type your study material here..."
    )
    
    if st.button("📥 Use Manual Text") and manual_text.strip():
        st.session_state.extracted_text = manual_text
        st.success(f"✅ Added {len(manual_text.split())} words")
    
    # Show extracted text preview
    if st.session_state.extracted_text:
        st.markdown("### 📄 Extracted Text Preview")
        preview_text = st.session_state.extracted_text[:500] + "..." if len(st.session_state.extracted_text) > 500 else st.session_state.extracted_text
        st.text_area("Preview", preview_text, height=150, disabled=True, label_visibility="collapsed")
        st.info(f"📊 Total words: {len(st.session_state.extracted_text.split())}")

def summary_page(summary_length, gemini_api_key):
    """Page for text summarization using Gemini AI"""
    st.header("📝 Smart Summary")
    
    # Check API key
    if not gemini_api_key or not GEMINI_AVAILABLE:
        st.warning("⚠️ Please configure your Gemini API key in the sidebar to use summarization features.")
        return
    
    # Check if text is available
    if not st.session_state.extracted_text:
        st.info("📄 Welcome to NoteWise! Please extract or enter text first using the **Text Extraction** page to get started with AI-powered summaries.")
        st.markdown(
            """
            <div style='text-align: center; margin: 2rem 0; padding: 2rem; background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%); border-radius: 15px; color: white;'>
                <h3 style='margin: 0 0 1rem 0;'>🧠 Ready to Transform Your Notes?</h3>
                <p style='margin: 0; font-size: 1.1em; opacity: 0.9;'>
                    Upload your documents and let AI create smart summaries, practice questions, and study materials!
                </p>
            </div>
            """, 
            unsafe_allow_html=True
        )
        return
    
    # Two-column layout
    col1, col2 = st.columns([1, 1])
    
    with col1:
        st.subheader("📝 Original Text")
        st.text_area("Original Text", st.session_state.extracted_text, height=300, disabled=True, label_visibility="collapsed")
    
    with col2:
        st.subheader(f"📄 {summary_length} Summary")
        
        # Summary options
        col2_1, col2_2 = st.columns(2)
        with col2_1:
            if st.button("🎯 Generate Summary", type="primary"):
                with st.spinner("Generating summary with Gemini AI..."):
                    try:
                        summarizer = GeminiTextSummarizer(api_key=gemini_api_key)
                        if not summarizer.is_available():
                            st.error("❌ Gemini API key is invalid or service unavailable")
                            return
                        
                        summary = summarizer.summarize(
                            st.session_state.extracted_text, 
                            length=summary_length.lower()
                        )
                        st.session_state.summary = summary
                        st.session_state.summary_type = summary_length.lower()
                        
                        # Save to storage if available
                        if STORAGE_AVAILABLE and hasattr(st.session_state, 'storage'):
                            try:
                                summary_id = st.session_state.storage.save_summary(
                                    st.session_state.session_id,
                                    st.session_state.extracted_text,
                                    summary,
                                    summary_length.lower()
                                )
                                st.session_state.last_summary_id = summary_id
                            except Exception as e:
                                print(f"Storage error: {e}")
                        
                        st.success("✅ Summary generated!")
                    except Exception as e:
                        st.error(f"❌ Error generating summary: {str(e)}")
        
        with col2_2:
            if st.button("📝 Bullet Points"):
                with st.spinner("Creating bullet-point summary..."):
                    try:
                        summarizer = GeminiTextSummarizer(api_key=gemini_api_key)
                        bullet_summary = summarizer.generate_bullet_summary(st.session_state.extracted_text, 5)
                        st.session_state.summary = bullet_summary
                        st.session_state.summary_type = "bullet_points"
                        
                        # Save to storage if available
                        if STORAGE_AVAILABLE and hasattr(st.session_state, 'storage'):
                            try:
                                summary_id = st.session_state.storage.save_summary(
                                    st.session_state.session_id,
                                    st.session_state.extracted_text,
                                    bullet_summary,
                                    "bullet_points"
                                )
                                st.session_state.last_summary_id = summary_id
                            except Exception as e:
                                print(f"Storage error: {e}")
                        
                        st.success("✅ Bullet summary generated!")
                    except Exception as e:
                        st.error(f"❌ Error generating bullet summary: {str(e)}")
        
        # Display summary with enhanced export options
        if st.session_state.summary:
            st.markdown("### Generated Summary:")
            st.write(st.session_state.summary)
            
            # Enhanced export options
            if EXPORT_AVAILABLE:
                st.markdown("#### 📤 Export Options")
                col_export1, col_export2, col_export3 = st.columns(3)
                
                with col_export1:
                    # Text export (enhanced)
                    exporter = EnhancedExporter()
                    enhanced_text = exporter.export_summary_text(
                        st.session_state.summary,
                        st.session_state.extracted_text,
                        getattr(st.session_state, 'summary_type', 'standard')
                    )
                    st.download_button(
                        label="📄 Download Text",
                        data=enhanced_text,
                        file_name="NoteWise_Summary.txt",
                        mime="text/plain"
                    )
                
                with col_export2:
                    # PDF export
                    try:
                        pdf_data = exporter.export_summary_pdf(
                            st.session_state.summary,
                            st.session_state.extracted_text,
                            getattr(st.session_state, 'summary_type', 'standard')
                        )
                        st.download_button(
                            label="📄 Download PDF",
                            data=pdf_data,
                            file_name="NoteWise_Summary.pdf",
                            mime="application/pdf"
                        )
                    except ImportError:
                        st.error("PDF export requires ReportLab library")
                
                with col_export3:
                    # Combined export (if questions exist)
                    if hasattr(st.session_state, 'questions') and st.session_state.questions:
                        try:
                            combined_pdf = exporter.create_combined_export_pdf(
                                st.session_state.summary,
                                st.session_state.questions,
                                st.session_state.extracted_text,
                                getattr(st.session_state, 'summary_type', 'standard')
                            )
                            st.download_button(
                                label="📚 Combined PDF",
                                data=combined_pdf,
                                file_name="NoteWise_StudyMaterial.pdf",
                                mime="application/pdf"
                            )
                        except ImportError:
                            st.error("PDF export requires ReportLab library")
                    else:
                        st.info("Generate questions to enable combined export")
            else:
                # Fallback to basic download
                st.download_button(
                    label="💾 Download Summary",
                    data=st.session_state.summary,
                    file_name="NoteWise_Summary.txt",
                    mime="text/plain"
                )

def question_generation_page(question_count, gemini_api_key):
    """Page for question generation using Gemini AI"""
    st.header("❓ Practice Questions")
    
    # Check API key
    if not gemini_api_key or not GEMINI_AVAILABLE:
        st.warning("⚠️ Please configure your Gemini API key in the sidebar to use question generation features.")
        return
    
    # Check if text is available
    if not st.session_state.extracted_text:
        st.info("📄 Welcome to NoteWise! Please extract or enter text first using the **Text Extraction** page to generate practice questions.")
        st.markdown(
            """
            <div style='text-align: center; margin: 2rem 0; padding: 2rem; background: linear-gradient(135deg, #667eea 0%, #764ba2 100%); border-radius: 15px; color: white;'>
                <h3 style='margin: 0 0 1rem 0;'>❓ Create Smart Practice Questions</h3>
                <p style='margin: 0; font-size: 1.1em; opacity: 0.9;'>
                    Turn your study material into MCQs, True/False, and Short Answer questions powered by AI!
                </p>
            </div>
            """, 
            unsafe_allow_html=True
        )
        return
    
    st.markdown(f"**Text Length:** {len(st.session_state.extracted_text.split())} words")
    
    # Question type selection
    st.subheader("🎯 Question Types")
    col1, col2, col3 = st.columns(3)
    
    with col1:
        mcq_enabled = st.checkbox("📝 Multiple Choice", value=True, key="mcq_checkbox")
    
    with col2:
        tf_enabled = st.checkbox("✅ True/False", value=True, key="tf_checkbox")
    
    with col3:
        short_enabled = st.checkbox("💬 Short Answer", value=True, key="short_checkbox")
    
    # Validate at least one type is selected
    selected_types = []
    if mcq_enabled:
        selected_types.append("multiple choice")
    if tf_enabled:
        selected_types.append("true/false")
    if short_enabled:
        selected_types.append("short answer")
    
    if not selected_types:
        st.warning("⚠️ Please select at least one question type.")
        return
    
    # Calculate total questions that will be generated
    total_questions = question_count * len(selected_types)
    
    if len(selected_types) == 1:
        st.info(f"📊 Will generate {question_count} {selected_types[0]} questions")
    else:
        st.info(f"📊 Will generate {question_count} questions for each type: {', '.join(selected_types)} (Total: {total_questions} questions)")
    
    # Generate questions
    if st.button("🎯 Generate Questions", type="primary"):
        with st.spinner(f"Generating {total_questions} questions with Gemini 2.0 Flash..."):
            try:
                generator = GeminiQuestionGenerator(api_key=gemini_api_key)
                questions = generator.generate_questions(
                    st.session_state.extracted_text,
                    question_types=selected_types,
                    count=question_count
                )
                st.session_state.questions = questions
                
                # Save to storage if available
                if STORAGE_AVAILABLE and hasattr(st.session_state, 'storage'):
                    try:
                        question_set_id = st.session_state.storage.save_questions(
                            st.session_state.session_id,
                            questions,
                            "mixed"
                        )
                        st.session_state.last_question_set_id = question_set_id
                    except Exception as e:
                        print(f"Storage error: {e}")
                
                st.success("✅ Questions generated!")
            except Exception as e:
                st.error(f"❌ Error generating questions: {str(e)}")
    
    # Display questions
    if st.session_state.questions:
        st.subheader("📝 Practice Questions")
        for i, question in enumerate(st.session_state.questions, 1):
            if 'error' not in question:
                with st.expander(f"Question {i}: {question.get('question', '')[:50]}..."):
                    st.markdown(f"**Question:** {question.get('question', '')}")
                    
                    # Display options for MCQ
                    if question.get('type') == 'Multiple Choice' and question.get('options'):
                        st.markdown("**Options:**")
                        for j, opt in enumerate(question.get('options', []), 1):
                            st.write(f"{j}. {opt}")
                    
                    # Display answer
                    st.success(f"**Answer:** {question.get('answer', '')}")
                    
                    # Display explanation if available
                    if question.get('explanation'):
                        st.info(f"**Explanation:** {question['explanation']}")
        
        # Generate dynamic filename based on question types in the session
        def get_question_filename(extension):
            if not hasattr(st.session_state, 'questions') or not st.session_state.questions:
                return f"NoteWise_Questions.{extension}"
            
            # Analyze question types in the current session
            question_types = set()
            for q in st.session_state.questions:
                if 'error' not in q and 'type' in q:
                    q_type = q['type'].lower()
                    if 'multiple choice' in q_type:
                        question_types.add('MCQ')
                    elif 'true/false' in q_type or 'true false' in q_type:
                        question_types.add('TrueFalse')
                    elif 'short answer' in q_type:
                        question_types.add('ShortAnswer')
            
            if question_types:
                # Sort for consistent naming
                type_names = sorted(list(question_types))
                filename = f"NoteWise_{'_'.join(type_names)}.{extension}"
                return filename
            else:
                return f"NoteWise_Questions.{extension}"
        
        # Enhanced download options for questions
        if EXPORT_AVAILABLE:
            st.markdown("#### 📤 Export Questions")
            col_q1, col_q2 = st.columns(2)
            
            with col_q1:
                # Enhanced text export
                exporter = EnhancedExporter()
                questions_text = exporter.export_questions_text(st.session_state.questions)
                st.download_button(
                    label="📄 Download Text",
                    data=questions_text,
                    file_name=get_question_filename("txt"),
                    mime="text/plain"
                )
            
            with col_q2:
                # PDF export
                try:
                    questions_pdf = exporter.export_questions_pdf(st.session_state.questions)
                    st.download_button(
                        label="📄 Download PDF",
                        data=questions_pdf,
                        file_name=get_question_filename("pdf"),
                        mime="application/pdf"
                    )
                except ImportError:
                    st.error("PDF export requires ReportLab library")
        else:
            # Basic text download
            questions_text = ""
            for i, q in enumerate(st.session_state.questions, 1):
                if 'error' not in q:
                    questions_text += f"Question {i}: {q.get('question', '')}\n"
                    if q.get('type') == 'Multiple Choice':
                        for j, opt in enumerate(q.get('options', []), 1):
                            questions_text += f"  {j}. {opt}\n"
                        questions_text += f"Answer: {q.get('answer', '')}\n"
                    else:
                        questions_text += f"Answer: {q.get('answer', '')}\n"
                    questions_text += f"Explanation: {q.get('explanation', '')}\n\n"
            
            st.download_button(
                label="💾 Download Questions",
                data=questions_text,
                file_name=get_question_filename("txt"),
                mime="text/plain"
            )

def history_page():
    """Page for viewing history and managing data"""
    st.header("📊 Session History")
    
    if st.session_state.extracted_text:
        st.subheader("📄 Current Session")
        st.write(f"**Text Length:** {len(st.session_state.extracted_text.split())} words")
        
        if st.session_state.summary:
            st.write("✅ Summary generated")
        else:
            st.write("⏳ No summary yet")
        
        if st.session_state.questions:
            st.write(f"✅ {len(st.session_state.questions)} questions generated")
        else:
            st.write("⏳ No questions yet")
    else:
        st.info("No current session. Upload and extract text to get started.")
    
    # Show storage history if available
    if STORAGE_AVAILABLE and hasattr(st.session_state, 'storage'):
        st.markdown("---")
        st.subheader("📚 Stored Content")
        
        # Get stored summaries
        try:
            summaries = st.session_state.storage.get_summaries(st.session_state.session_id)
            if summaries:
                st.write(f"**📄 Summaries:** {len(summaries)} stored")
                with st.expander("View Stored Summaries"):
                    for i, summary in enumerate(reversed(summaries[-5:]), 1):  # Show last 5
                        st.write(f"**{i}.** {summary['summary_type'].title()} - {summary['created_at']}")
                        st.write(f"Word count: {summary['word_count']}")
                        if st.button(f"📥 Load Summary {i}", key=f"load_summary_{summary['summary_id']}"):
                            st.session_state.summary = summary['summary']
                            st.rerun()
            
            # Get stored questions
            question_sets = st.session_state.storage.get_questions(st.session_state.session_id)
            if question_sets:
                st.write(f"**❓ Question Sets:** {len(question_sets)} stored")
                with st.expander("View Stored Questions"):
                    for i, q_set in enumerate(reversed(question_sets[-5:]), 1):  # Show last 5
                        st.write(f"**{i}.** {q_set['question_type'].title()} - {q_set['created_at']}")
                        st.write(f"Questions: {q_set['question_count']}")
                        if st.button(f"📥 Load Questions {i}", key=f"load_questions_{q_set['question_set_id']}"):
                            st.session_state.questions = q_set['questions']
                            st.rerun()
        except Exception as e:
            st.error(f"Error loading storage history: {e}")
    
    st.markdown("---")
    
    if st.button("🗑️ Clear Session", type="secondary"):
        st.session_state.extracted_text = ""
        st.session_state.summary = ""
        st.session_state.questions = []
        st.success("✅ Session cleared!")
        st.rerun()
    
    # Project landmark/footer
    st.markdown("---")
    st.markdown(
        """
        <div style='text-align: center; color: #666; font-size: 0.85em; margin-top: 2rem; padding: 1rem 0; border-top: 1px solid #eee;'>
            <p style='margin: 0;'>
                🧠 <strong>NoteWise</strong> - Smart Study Assistant | 
                <em>Transform your notes into interactive study material with AI</em>
            </p>
            <p style='margin: 0.5rem 0 0 0; font-size: 0.75em; opacity: 0.8;'>
                Powered by Google Gemini 2.0 Flash • Built with ❤️ for learners
            </p>
        </div>
        """, 
        unsafe_allow_html=True
    )

if __name__ == "__main__":
    main()
